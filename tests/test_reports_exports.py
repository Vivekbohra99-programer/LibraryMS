"""Tests for reports and CSV/Excel exports."""
from app.services.export import to_csv, to_excel


def test_export_to_csv():
    headers = ['ID', 'Title', 'Author']
    rows = [
        [1, 'Clean Code', 'Robert C. Martin'],
        [2, 'The Pragmatic Programmer', 'Andrew Hunt'],
    ]
    csv_str = to_csv(headers, rows)
    assert 'ID,Title,Author' in csv_str
    assert 'Clean Code' in csv_str
    assert 'Robert C. Martin' in csv_str


def test_export_to_excel():
    headers = ['ID', 'Title', 'Author']
    rows = [
        [1, 'Clean Code', 'Robert C. Martin'],
    ]
    excel_bytes = to_excel(headers, rows, sheet_name='Books')
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 0


def test_books_export_csv_endpoint(librarian_client):
    response = librarian_client.get('/reports/export/books?format=csv')
    assert response.status_code == 200
    assert response.mimetype == 'text/csv'
    assert b'Patterns of Enterprise Application Architecture' in response.data


def test_books_export_excel_endpoint(librarian_client):
    response = librarian_client.get('/reports/export/books?format=excel')
    assert response.status_code == 200
    assert 'spreadsheet' in response.mimetype


def test_members_export_csv_endpoint(librarian_client):
    response = librarian_client.get('/reports/export/members?format=csv')
    assert response.status_code == 200
    assert response.mimetype == 'text/csv'
    assert b'Test Member' in response.data


def test_reports_index_view(librarian_client):
    response = librarian_client.get('/reports')
    assert response.status_code == 200
    assert b'Reports' in response.data


# ---------- exports must not carry formulas ----------

def test_csv_export_neutralises_formula_text():
    import csv
    import io
    risky = ['=HYPERLINK("http://example.test","click")', '+cmd|x', '-2+3', '@SUM(A1:A9)', '\tTabbed']
    safe = ['Clean Code', '2026-10-11', '+91 98765-43210', '(555) 123-4567']
    out = list(csv.reader(io.StringIO(to_csv(['Text'], [[v] for v in risky + safe + [-5, 3.5]]))))[1:]
    cells = [r[0] for r in out]
    assert all(c.startswith("'") for c in cells[:5])                  # every risky cell is defused
    assert cells[0] == "'" + risky[0]                                  # and the original text is kept after the apostrophe
    assert cells[5:9] == safe                                          # titles, dates and phone numbers are untouched
    assert cells[9:] == ['-5', '3.5']                                  # real numbers are not text, so they are untouched


def test_excel_export_stores_formula_text_as_text():
    import io
    from openpyxl import load_workbook
    data = to_excel(['Text'], [['=1+1'], ['@SUM(1,2)'], ['Clean Code'], [42]])
    ws = load_workbook(io.BytesIO(data)).active
    assert (ws['A2'].value, ws['A2'].data_type) == ('=1+1', 's')      # text, not a formula
    assert ws['A3'].value == '@SUM(1,2)' and ws['A4'].value == 'Clean Code' and ws['A5'].value == 42


def test_books_export_defuses_a_formula_title(librarian_client, app):
    from app.extensions import db
    from app.models.book import Book
    with app.app_context():
        b = Book.query.first()
        b.title = '=2+2'
        db.session.commit()
    body = librarian_client.get('/reports/export/books?format=csv').get_data(as_text=True)
    assert "'=2+2" in body and ',=2+2,' not in body
